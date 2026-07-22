#!/usr/bin/env python3
"""Simple local PDF toolbox: tables to Excel, split, sign, merge, extract text.

Run it with the .xlsxenv Python (the "PDF Tool.command" launcher does this for you).
Everything works on files sitting on your computer. Nothing is uploaded anywhere.
"""

import os
import sys

import fitz  # PyMuPDF
from openpyxl import Workbook

# ---------------------------------------------------------------------------
# Small readable helpers (built for easy typing: drag a file in, press Enter)
# ---------------------------------------------------------------------------

DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")


def clean_path(raw):
    """Turn whatever the user typed/dragged into a real path.

    macOS Terminal drops a dragged file in as a path with backslash-escaped
    spaces and sometimes surrounding quotes. Strip all of that.
    """
    p = raw.strip()
    if len(p) >= 2 and p[0] == p[-1] and p[0] in "'\"":
        p = p[1:-1]
    p = p.replace("\\ ", " ").replace("\\'", "'").replace('\\"', '"')
    return os.path.expanduser(p.strip())


def ask(prompt):
    return input(prompt).strip()


def ask_pdf(prompt="  Drag your PDF here, then press Enter: "):
    while True:
        p = clean_path(ask(prompt))
        if not p:
            return None
        if not os.path.exists(p):
            print(f"  I can't find that file. Try again (or press Enter to go back).")
            continue
        if not p.lower().endswith(".pdf"):
            print("  That doesn't look like a PDF. Try again.")
            continue
        return p


def ask_int(prompt, lo=None, hi=None):
    while True:
        v = ask(prompt)
        if v == "":
            return None
        try:
            n = int(v)
        except ValueError:
            print("  Please type a whole number.")
            continue
        if lo is not None and n < lo:
            print(f"  Please type {lo} or higher.")
            continue
        if hi is not None and n > hi:
            print(f"  Please type {hi} or lower.")
            continue
        return n


def out_path(src, suffix, ext):
    """Build an output filename next to the source, e.g. report_tables.xlsx."""
    base = os.path.splitext(os.path.basename(src))[0]
    folder = os.path.dirname(src) or DESKTOP
    return os.path.join(folder, f"{base}{suffix}{ext}")


def done(path):
    print()
    print(f"  Done. Saved to:")
    print(f"  {path}")
    print()


def pause():
    input("  Press Enter to return to the menu. ")


# ---------------------------------------------------------------------------
# 1. Tables -> Excel
# ---------------------------------------------------------------------------

def tables_to_excel():
    print("\n== Turn PDF tables into an Excel file ==")
    src = ask_pdf()
    if not src:
        return
    doc = fitz.open(src)
    wb = Workbook()
    wb.remove(wb.active)  # start with no sheets
    found = 0
    for pno, page in enumerate(doc, start=1):
        tables = page.find_tables()
        for tno, table in enumerate(tables.tables, start=1):
            found += 1
            title = f"p{pno}_t{tno}"[:31]  # Excel sheet-name limit
            ws = wb.create_sheet(title=title)
            for row in table.extract():
                ws.append([("" if c is None else str(c)) for c in row])
    doc.close()
    if found == 0:
        print("\n  I didn't find any grid-style tables in that PDF.")
        print("  (This works on real tables with lines/columns, not scanned images.)")
        print()
        pause()
        return
    dst = out_path(src, "_tables", ".xlsx")
    wb.save(dst)
    print(f"\n  Found {found} table(s).")
    done(dst)
    pause()


# ---------------------------------------------------------------------------
# 2. Split a long PDF into pieces
# ---------------------------------------------------------------------------

def split_pdf():
    print("\n== Cut a long PDF into smaller pieces ==")
    src = ask_pdf()
    if not src:
        return
    doc = fitz.open(src)
    total = doc.page_count
    print(f"\n  That PDF has {total} pages.")
    print("  How do you want to cut it?")
    print("    1) Every N pages (e.g. one file per 10 pages)")
    print("    2) One specific range of pages (e.g. pages 5 to 12)")
    choice = ask("  Type 1 or 2: ")

    saved = []
    if choice == "1":
        n = ask_int(f"  Pages per piece (1-{total}): ", lo=1, hi=total)
        if not n:
            doc.close()
            return
        part = 1
        for start in range(0, total, n):
            end = min(start + n, total)
            piece = fitz.open()
            piece.insert_pdf(doc, from_page=start, to_page=end - 1)
            dst = out_path(src, f"_part{part}", ".pdf")
            piece.save(dst)
            piece.close()
            saved.append(dst)
            part += 1
    elif choice == "2":
        a = ask_int(f"  First page (1-{total}): ", lo=1, hi=total)
        if not a:
            doc.close()
            return
        b = ask_int(f"  Last page ({a}-{total}): ", lo=a, hi=total)
        if not b:
            doc.close()
            return
        piece = fitz.open()
        piece.insert_pdf(doc, from_page=a - 1, to_page=b - 1)
        dst = out_path(src, f"_pages{a}-{b}", ".pdf")
        piece.save(dst)
        piece.close()
        saved.append(dst)
    else:
        print("  Cancelled.")
        doc.close()
        return

    doc.close()
    print(f"\n  Made {len(saved)} file(s):")
    for s in saved:
        print(f"  {s}")
    print()
    pause()


# ---------------------------------------------------------------------------
# 3. Add a signature image onto a page
# ---------------------------------------------------------------------------

CORNERS = {
    "1": ("top-left", 0.0, 0.0),
    "2": ("top-right", 1.0, 0.0),
    "3": ("bottom-left", 0.0, 1.0),
    "4": ("bottom-right", 1.0, 1.0),
    "5": ("center", 0.5, 0.5),
}


def add_signature():
    print("\n== Add your signature to a PDF ==")
    print("  You'll need a picture of your signature (PNG with a see-through")
    print("  background works best; a JPG is fine too).")
    src = ask_pdf("  Drag the PDF here, then press Enter: ")
    if not src:
        return
    while True:
        sig = clean_path(ask("  Drag your signature image here, then press Enter: "))
        if not sig:
            return
        if os.path.exists(sig) and sig.lower().endswith((".png", ".jpg", ".jpeg")):
            break
        print("  I need a .png or .jpg image that exists. Try again.")

    doc = fitz.open(src)
    total = doc.page_count
    pno = ask_int(f"\n  Which page to sign? (1-{total}): ", lo=1, hi=total)
    if not pno:
        doc.close()
        return

    print("\n  Where on the page?")
    for k, (name, _, _) in CORNERS.items():
        print(f"    {k}) {name}")
    pos = ask("  Type 1-5: ")
    if pos not in CORNERS:
        print("  Cancelled.")
        doc.close()
        return
    _, fx, fy = CORNERS[pos]

    width = ask_int("  Signature width in points (Enter for 150): ", lo=20, hi=600)
    if not width:
        width = 150

    # Keep the image's real proportions so the signature isn't squished.
    img = fitz.open(sig)
    r = img[0].rect
    img.close()
    ratio = (r.height / r.width) if r.width else 0.4
    height = width * ratio

    page = doc[pno - 1]
    pw, ph = page.rect.width, page.rect.height
    margin = 36  # half inch
    x = margin + fx * (pw - width - 2 * margin)
    y = margin + fy * (ph - height - 2 * margin)
    rect = fitz.Rect(x, y, x + width, y + height)
    page.insert_image(rect, filename=sig)

    dst = out_path(src, "_signed", ".pdf")
    doc.save(dst)
    doc.close()
    done(dst)
    pause()


# ---------------------------------------------------------------------------
# 4. Merge several PDFs into one
# ---------------------------------------------------------------------------

def merge_pdfs():
    print("\n== Combine several PDFs into one ==")
    print("  Add files one at a time. Press Enter on an empty line when done.")
    files = []
    while True:
        p = ask_pdf(f"  Drag PDF #{len(files) + 1} here (or Enter to finish): ")
        if not p:
            break
        files.append(p)
        print(f"    added: {os.path.basename(p)}")
    if len(files) < 2:
        print("\n  Need at least 2 PDFs to merge. Cancelled.")
        pause()
        return
    out = fitz.open()
    for f in files:
        d = fitz.open(f)
        out.insert_pdf(d)
        d.close()
    dst = os.path.join(os.path.dirname(files[0]) or DESKTOP, "merged.pdf")
    out.save(dst)
    out.close()
    done(dst)
    pause()


# ---------------------------------------------------------------------------
# 5. Pull the text out of a PDF
# ---------------------------------------------------------------------------

def extract_text():
    print("\n== Save all the text from a PDF as a plain text file ==")
    src = ask_pdf()
    if not src:
        return
    doc = fitz.open(src)
    parts = []
    for pno, page in enumerate(doc, start=1):
        parts.append(f"----- Page {pno} -----\n{page.get_text()}")
    doc.close()
    dst = out_path(src, "_text", ".txt")
    with open(dst, "w", encoding="utf-8") as fh:
        fh.write("\n\n".join(parts))
    done(dst)
    pause()


# ---------------------------------------------------------------------------
# Menu
# ---------------------------------------------------------------------------

MENU = [
    ("1", "Turn PDF tables into an Excel file", tables_to_excel),
    ("2", "Cut a long PDF into smaller pieces", split_pdf),
    ("3", "Add my signature to a PDF", add_signature),
    ("4", "Combine several PDFs into one", merge_pdfs),
    ("5", "Save a PDF's text as a text file", extract_text),
]


def main():
    while True:
        print("\n" + "=" * 44)
        print("        PDF TOOL")
        print("=" * 44)
        for key, label, _ in MENU:
            print(f"   {key})  {label}")
        print("   q)  Quit")
        print("=" * 44)
        choice = ask("Type a number and press Enter: ").lower()
        if choice in ("q", "quit", "exit"):
            print("Bye.")
            return
        for key, _, fn in MENU:
            if choice == key:
                try:
                    fn()
                except Exception as e:  # keep the tool alive on any single error
                    print(f"\n  Something went wrong: {e}")
                    pause()
                break
        else:
            print("  I didn't recognize that. Type one of the numbers, or q to quit.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nBye.")
        sys.exit(0)
